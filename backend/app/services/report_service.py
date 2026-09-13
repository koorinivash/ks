import calendar
import io
from datetime import datetime, timezone
from typing import Any, Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

NAVY = "1B2A4A"


def _month_label(month: int, year: int) -> str:
    return f"{calendar.month_name[month]} {year}"


def _row_date_label(row: dict[str, Any]) -> str:
    payment_date = row.get("payment_date")
    if payment_date:
        return payment_date.strftime("%d %b %Y")
    month = row.get("month")
    year = row.get("year")
    if month and year:
        return f"{calendar.month_abbr[month]} {year}"
    return "-"


def generate_excel_report(
    rows: list[dict[str, Any]],
    summary: dict[str, Any],
    month: Optional[int],
    year: Optional[int],
) -> bytes:
    is_full = month is None or year is None
    wb = Workbook()
    ws = wb.active
    ws.title = "Full Report" if is_full else "Monthly Report"

    bold = Font(bold=True)
    title_font = Font(bold=True, size=14, color=NAVY)
    header_fill = PatternFill(start_color=NAVY, end_color=NAVY, fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")

    last_col = "D"
    ws.merge_cells(f"A1:{last_col}1")
    ws["A1"] = f"KS - {'Full Report (All Time)' if is_full else 'Monthly Report'}"
    ws["A1"].font = title_font
    ws["A1"].alignment = Alignment(horizontal="center")

    ws.merge_cells(f"A2:{last_col}2")
    ws["A2"] = "Period: All Time" if is_full else f"Month: {_month_label(month, year)}"
    ws["A2"].font = bold
    ws["A2"].alignment = Alignment(horizontal="center")

    headers = ["S.No", "Name", "Date", "Amount"]
    header_row = 4
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    row_idx = header_row + 1
    for i, row in enumerate(rows, start=1):
        ws.cell(row=row_idx, column=1, value=i)
        ws.cell(row=row_idx, column=2, value=row["name"])
        ws.cell(row=row_idx, column=3, value=_row_date_label(row))
        ws.cell(row=row_idx, column=4, value=row["paid"])
        row_idx += 1

    totals_row = row_idx + 1
    ws.cell(row=totals_row, column=2, value="Total Collected:").font = bold
    ws.cell(row=totals_row, column=4, value=summary["collected_amount"]).font = bold

    generated_row = totals_row + 3
    ws.cell(
        row=generated_row,
        column=1,
        value=f"Generated on {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')}",
    )

    widths = [8, 26, 16, 16]
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def generate_pdf_report(
    rows: list[dict[str, Any]],
    summary: dict[str, Any],
    month: Optional[int],
    year: Optional[int],
) -> bytes:
    is_full = month is None or year is None
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        topMargin=18 * mm,
        bottomMargin=18 * mm,
        leftMargin=15 * mm,
        rightMargin=15 * mm,
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "KSTitle", parent=styles["Title"], textColor=colors.HexColor(f"#{NAVY}"), fontSize=20
    )
    subtitle_style = ParagraphStyle(
        "KSSubtitle", parent=styles["Normal"], alignment=1, fontSize=12, textColor=colors.HexColor(f"#{NAVY}")
    )
    month_style = ParagraphStyle(
        "KSMonth", parent=styles["Normal"], alignment=1, fontSize=14, spaceAfter=12
    )

    elements = [
        Paragraph("KS", title_style),
        Paragraph("Full Finance Report" if is_full else "Monthly Finance Report", subtitle_style),
        Spacer(1, 6),
        Paragraph("All Time" if is_full else _month_label(month, year), month_style),
    ]

    summary_data = [
        ["Total People" if not is_full else "Total Records", str(summary["total_people"])],
        ["Total Collected", f"Rs. {summary['collected_amount']:,.2f}"],
    ]
    summary_table = Table(summary_data, colWidths=[80 * mm, 80 * mm])
    summary_table.setStyle(
        TableStyle(
            [
                ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
                ("LINEBELOW", (0, 0), (-1, -1), 0.5, colors.HexColor("#E5E7EB")),
            ]
        )
    )
    elements.append(summary_table)
    elements.append(Spacer(1, 16))

    table_data = [["S.No", "Name", "Date", "Amount"]]
    col_widths = [20 * mm, 60 * mm, 40 * mm, 40 * mm]
    for i, row in enumerate(rows, start=1):
        table_data.append([str(i), row["name"], _row_date_label(row), f"Rs. {row['paid']:,.2f}"])

    detail_table = Table(table_data, colWidths=col_widths, repeatRows=1)
    style_commands = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(f"#{NAVY}")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#E5E7EB")),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F9FAFB")]),
    ]
    detail_table.setStyle(TableStyle(style_commands))
    elements.append(detail_table)
    elements.append(Spacer(1, 20))

    footer_style = ParagraphStyle("KSFooter", parent=styles["Normal"], alignment=1, fontSize=9, textColor=colors.grey)
    elements.append(Paragraph("Generated by KS", footer_style))
    elements.append(
        Paragraph(
            f"Generated on {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')}",
            footer_style,
        )
    )

    doc.build(elements)
    return buffer.getvalue()
