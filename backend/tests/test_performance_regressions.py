import asyncio
import io
import threading
from decimal import localcontext
from unittest.mock import patch

import pytest
from openpyxl import load_workbook

from app.services.finance_service import get_monthly_report_rows, summarize_rows


async def make_person(client, name, amount=1000, count=1, status="active"):
    response = await client.post("/api/people", json={
        "name": name, "default_monthly_amount": amount, "count": count, "status": status,
    })
    assert response.status_code == 201
    return response.json()["id"]


async def pay(client, person_id, month=9, amount=1000, paid=400):
    response = await client.post("/api/monthly-records", json={
        "person_id": person_id, "year": 2026, "month": month,
        "amount": amount, "paid_amount": paid, "payment_date": "2026-09-13T12:00:00Z",
    })
    assert response.status_code == 201
    return response.json()["id"]


async def test_dashboard_matches_monthly_report_and_crud(client, mock_database):
    ids = [await make_person(client, f"Person {i}", count=i + 1) for i in range(4)]
    inactive = await make_person(client, "Inactive", status="inactive")
    await pay(client, inactive, paid=9000)
    await pay(client, ids[0], amount=1000, paid=0)
    record = await pay(client, ids[1], amount=2000, paid=500)
    await pay(client, ids[2], amount=1000, paid=1500)
    await pay(client, ids[2], month=8, paid=3000)

    async def check():
        reference = summarize_rows(await get_monthly_report_rows(mock_database, 9, 2026))
        response = await client.get("/api/dashboard?month=9&year=2026")
        assert response.status_code == 200
        for key, expected in reference.items():
            assert response.json()[key] == pytest.approx(expected)
        return response.json()

    assert (await check())["expected_amount"] == 8000
    await client.put(f"/api/monthly-records/{record}", json={"paid_amount": 1500})
    assert (await check())["collected_amount"] == 3000
    await client.delete(f"/api/monthly-records/{record}")
    await check()
    await client.put(f"/api/people/{ids[3]}", json={"count": 5})
    await check()
    await client.delete(f"/api/people/{ids[2]}")
    await check()


async def test_empty_and_inactive_dashboard(client):
    assert (await client.get("/api/dashboard")).json()["has_any_people"] is False
    await make_person(client, "Inactive", status="inactive")
    data = (await client.get("/api/dashboard")).json()
    assert data["has_any_people"] is True
    assert data["total_people"] == data["expected_amount"] == data["collected_amount"] == 0


async def test_payment_picker_returns_only_active_names_and_ids(client):
    person = await make_person(client, "Active")
    await make_person(client, "Inactive", status="inactive")
    response = await client.get("/api/people/options")
    assert response.status_code == 200
    assert response.json() == [{"id": person, "name": "Active"}]


async def test_dashboard_preserves_small_override_of_large_default(client):
    person = await make_person(client, "Large default", amount=10**16)
    await pay(client, person, amount=1, paid=0.25)
    # mongomock implements $toDecimal with Python's context; match Decimal128 precision.
    with localcontext() as context:
        context.prec = 34
        response = await client.get("/api/dashboard?month=9&year=2026")
    assert response.status_code == 200
    assert response.json()["expected_amount"] == 1
    assert response.json()["pending_amount"] == 0.75


async def test_people_stats_pagination_and_literal_search(client, mock_database):
    ids = [await make_person(client, name) for name in ["A.*", "Aaron", "Bala"]]
    await pay(client, ids[0], paid=400)
    await pay(client, ids[0], month=8, paid=1000)
    collection_type = type(mock_database["monthly_records"])
    original = collection_type.aggregate
    calls = []

    def counted(self, *args, **kwargs):
        calls.append(1)
        return original(self, *args, **kwargs)

    with patch.object(collection_type, "aggregate", counted):
        data = (await client.get("/api/people?month=9&year=2026")).json()
    assert len(calls) == 1
    assert data[0]["total_contributed"] == 1400
    assert data[0]["paid_months"] == 1
    assert data[0]["pending_months"] == 1
    assert data[0]["current_month_paid"] == 400
    assert data[0]["current_month_status"] == "partial"
    assert data[1]["current_month_status"] == "pending"
    page1 = (await client.get("/api/people?page=1&limit=2")).json()
    page2 = (await client.get("/api/people?page=2&limit=2")).json()
    assert [p["id"] for p in page1 + page2] == [p["id"] for p in data]
    assert len((await client.get("/api/people", params={"search": ".*"})).json()) == 1
    assert (await client.get("/api/people?page=0")).status_code == 422


@pytest.mark.parametrize("period", ["monthly", "full"])
@pytest.mark.parametrize("kind", ["pdf", "excel"])
async def test_exports(client, period, kind):
    person = await make_person(client, "Export Test")
    await pay(client, person)
    response = await client.get(f"/api/reports/{period}/{kind}?month=9&year=2026")
    assert response.status_code == 200
    if kind == "pdf":
        assert response.content.startswith(b"%PDF")
    else:
        workbook = load_workbook(io.BytesIO(response.content))
        assert any("Export Test" in row for row in workbook.active.iter_rows(values_only=True))


async def test_report_renderer_runs_off_event_loop(client):
    from app.routes import reports
    event_loop_thread = threading.get_ident()
    started, release = threading.Event(), threading.Event()

    def render(*args):
        assert threading.get_ident() != event_loop_thread
        started.set()
        assert release.wait(timeout=3)
        return b"%PDF-test"

    with patch.object(reports, "_render_report", render):
        task = asyncio.create_task(client.get("/api/reports/full/pdf"))
        try:
            assert await asyncio.to_thread(started.wait, 3)
            response = await asyncio.wait_for(client.get("/health"), timeout=1)
            assert response.status_code == 200
            assert "app;dur=" in response.headers["server-timing"]
        finally:
            release.set()
        assert (await task).status_code == 200
